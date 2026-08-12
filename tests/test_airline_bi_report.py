from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from app.core.bi.airline_report import is_airline_operations_dataset
from app.core.bi.report_service import build_bi_report
from app.core.reporting.bi_exports import validate_powerbi_project_package, validate_tableau_packaged_workbook


def _flight_frame() -> pd.DataFrame:
    return pd.DataFrame({
        "FlightDate": ["2026-05-01", "2026-05-01", "2026-05-02", "2026-05-02", "2026-05-03", "2026-05-03"],
        "IATA_CODE_Reporting_Airline": ["AA", "AA", "DL", "DL", "UA", "UA"],
        "Origin": ["JFK", "JFK", "ATL", "ATL", "ORD", "ORD"],
        "Dest": ["LAX", "SFO", "LAX", "JFK", "SFO", "LAX"],
        "Cancelled": [0, 1, 0, 0, 0, 0],
        "Diverted": [0, 0, 0, 0, 1, 0],
        "ArrDelay": [10, None, 45, -5, None, 20],
        "DepDelay": [5, None, 35, -10, 50, 15],
        "CancellationCode": [None, "B", None, None, None, None],
        "CarrierDelay": [0, 0, 20, 0, 0, 10],
        "WeatherDelay": [0, 0, 0, 0, 0, 0],
        "NASDelay": [0, 0, 10, 0, 0, 5],
        "SecurityDelay": [0, 0, 0, 0, 0, 0],
        "LateAircraftDelay": [0, 0, 15, 0, 0, 5],
        "Distance": [2475, 2586, 1946, 760, 1846, 1744],
    })


def test_airline_report_has_source_backed_metrics_and_all_client_formats():
    frame = _flight_frame()
    assert is_airline_operations_dataset(frame) is True
    result = build_bi_report(frame, dataset_name="BTS May 2026", source_version_id="bts-test")
    metrics = {item["label"]: item["value"] for item in result["kpis"]}
    assert metrics["Scheduled Flights"] == 6
    assert metrics["On-Time Arrival %"] == 50.0
    assert round(metrics["Cancellation Rate"], 6) == round(100 / 6, 6)
    assert metrics["Average Arrival Delay"] == 17.5
    assert metrics["Total Positive Delay"] == 75.0
    assert result["source"]["domain"] == "airline_operations"
    assert result["source"]["grain"] == "one scheduled flight record"
    assert {item["source_ref"] for item in result["charts"]} == {"chart:daily_flights", "chart:carrier_ontime", "chart:origin_delay", "chart:delay_causes"}
    assert validate_powerbi_project_package(result["powerbi_bytes"])["valid"] is True
    assert validate_tableau_packaged_workbook(result["tableau_twbx_bytes"])["valid"] is True
    workbook = load_workbook(BytesIO(result["xlsx_bytes"]), read_only=False)
    assert "Dashboard" in workbook.sheetnames
    with ZipFile(BytesIO(result["powerbi_bytes"])) as package:
        assert "ENTERPRISE_ARCHITECTURE.md" in package.namelist()


def test_airline_report_accepts_curated_explicit_unit_names():
    curated = _flight_frame().rename(columns={
        "ArrDelay": "ArrivalDelayMinutes",
        "DepDelay": "DepartureDelayMinutes",
        "AirTime": "AirTimeMinutes",
        "Distance": "DistanceMiles",
    })
    result = build_bi_report(curated, dataset_name="Curated BTS")
    metrics = {item["label"]: item["value"] for item in result["kpis"]}
    assert metrics["On-Time Arrival %"] == 50.0
    assert metrics["Average Arrival Delay"] == 17.5


def test_airline_api_reuses_version_aware_report_cache(client):
    payload = _flight_frame().to_csv(index=False).encode("utf-8")
    imported = client.post("/api/v1/datasets/import", files={"file": ("bts.csv", payload, "text/csv")})
    assert imported.status_code == 200
    dataset_id = imported.json()["dataset_id"]
    first = client.get(f"/api/v1/datasets/{dataset_id}/bi_report", params={"template_id": "executive"})
    second = client.get(f"/api/v1/datasets/{dataset_id}/bi_report", params={"template_id": "executive"})
    assert first.status_code == second.status_code == 200
    assert first.json()["report_id"] == second.json()["report_id"]
    filtered = client.get(f"/api/v1/datasets/{dataset_id}/bi_report", params={"template_id": "executive", "carrier": "AA"})
    assert filtered.status_code == 200
    assert filtered.json()["report_id"] != first.json()["report_id"]
    assert filtered.json()["source"]["row_count"] == 2

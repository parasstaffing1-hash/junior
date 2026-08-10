import pytest
from app.core.reporting.xlsx_generator import build_xlsx_bytes, ExcelGenerationError
from openpyxl import load_workbook
from io import BytesIO

def test_excel_generator_basic():
    manifest = {
        "report": {
            "title": "Test Report",
            "description": "Test Description",
            "parameters": {"date": "2026-01-01"}
        },
        "sections": [
            {
                "section_type": "kpi",
                "title": "Total Users",
                "content": {"value": 1500, "formatted_value": "1,500", "percent_change": 5.5, "target_status": "Met"}
            },
            {
                "section_type": "text",
                "title": "Observation",
                "content": "=1+1"  # Should be escaped
            },
            {
                "section_type": "table",
                "title": "User Data",
                "content": {
                    "columns": ["ID", "Name"],
                    "rows": [{"ID": 1, "Name": "Alice"}, {"ID": 2, "Name": "Bob"}]
                }
            }
        ]
    }
    
    excel_bytes = build_xlsx_bytes(manifest)
    assert excel_bytes is not None
    assert len(excel_bytes) > 0

    wb = load_workbook(BytesIO(excel_bytes))
    assert "Executive Summary" in wb.sheetnames
    assert "KPIs" in wb.sheetnames
    assert "Findings" in wb.sheetnames
    
    ws = wb["Findings"]
    assert ws["A2"].value == "Observation"
    assert ws["B2"].value == "'=1+1"  # Prevented formula injection

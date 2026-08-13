from __future__ import annotations

from io import BytesIO
import json
import sqlite3
from zipfile import ZipFile

from openpyxl import load_workbook
import pandas as pd

from app.core.eda.report import generate_eda_report
from app.core.eda.relationships import analyze_categorical_relationships
from app.core.professional.analysis import build_business_analysis
from app.core.projects.fixtures import build_project_fixture
from app.core.projects.workbench import build_project
from app.core.reporting.bi_exports import validate_powerbi_project_package
from app.core.sql.workbench import run_sql_proficiency_benchmark


def test_large_high_cardinality_eda_is_bounded_and_explainable():
    rows = 6000
    frame = pd.DataFrame(
        {
            "record_id": [f"RID-{index:05d}" for index in range(rows)],
            "customer_ref": [f"CUS-{index:05d}" for index in range(rows)],
            "region": ["North", "South", "East", "West"] * (rows // 4),
            "channel": ["Online", "Store"] * (rows // 2),
            "revenue": [float((index % 200) + 1) for index in range(rows)],
            "profit": [float((index % 80) - 10) for index in range(rows)],
        }
    )

    report = generate_eda_report(frame)
    basis = report["analysis_basis"]
    assert basis["relationship_sampled"] is True
    assert basis["relationship_rows"] == 5000
    assert {"record_id", "customer_ref"}.issubset(basis["skipped_high_cardinality_relationship_columns"])
    assert len(json.dumps(report)) < 1_000_000

    guarded = analyze_categorical_relationships(frame, left="record_id", right="customer_ref")
    assert guarded["relationships"][0]["status"] == "too_many_levels"
    assert guarded["skipped_pair_count"] == 1


def test_sql_professional_benchmark_and_uploaded_database(client, tmp_path):
    benchmark = run_sql_proficiency_benchmark()
    assert benchmark["status"] == "PASS"
    assert benchmark["score"] >= 80
    assert benchmark["passed"] == 20
    assert benchmark["question_count"] == 20
    assert not [item for item in benchmark["results"] if not item["passed"]]

    api_benchmark = client.post("/api/v1/sql/proficiency-benchmark")
    assert api_benchmark.status_code == 200
    assert api_benchmark.json()["score"] == benchmark["score"]

    database_path = tmp_path / "business.sqlite"
    connection = sqlite3.connect(database_path)
    connection.executescript(
        """
        CREATE TABLE customers (customer_id INTEGER PRIMARY KEY, region TEXT);
        CREATE TABLE orders (order_id INTEGER PRIMARY KEY, customer_id INTEGER, order_date TEXT, revenue REAL);
        INSERT INTO customers VALUES (1, 'North'), (2, 'South');
        INSERT INTO orders VALUES (1, 1, '2025-01-01', 100.0), (2, 1, '2025-02-01', 150.0), (3, 2, '2025-02-03', 90.0);
        """
    )
    connection.close()
    sql = """
        WITH monthly AS (
            SELECT c.region, strftime('%Y-%m', o.order_date) AS month, SUM(o.revenue) AS revenue
            FROM orders o JOIN customers c ON c.customer_id=o.customer_id
            GROUP BY c.region, month
        )
        SELECT region, month, revenue,
               DENSE_RANK() OVER (PARTITION BY month ORDER BY revenue DESC) AS revenue_rank,
               SUM(revenue) OVER (PARTITION BY region ORDER BY month) AS running_revenue
        FROM monthly
    """
    response = client.post(
        "/api/v1/sql/database-query",
        files={"file": ("business.sqlite", database_path.read_bytes(), "application/vnd.sqlite3")},
        data={"sql": sql, "max_rows": "100", "timeout_seconds": "5"},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert {"customers", "orders"} == {item["name"] for item in result["tables"]}
    assert {"cte", "join", "group_by", "date_manipulation", "window_function", "ranking", "running_total"}.issubset(result["features"])
    assert result["optimization"]["query_plan"]


def test_professional_excel_and_powerbi_artifacts():
    frame = build_project_fixture("superstore_sales_dashboard", rows=48)
    result = build_project(frame, project_id="superstore_sales_dashboard")

    workbook = load_workbook(BytesIO(result["xlsx_bytes"]), data_only=False)
    required_sheets = {"Dashboard", "Raw Data", "Pivot Summary", "Formula Lab", "Power Query", "Workbook Guide", "Exceptions", "Reconciliation", "MIS Control", "MIS Automation"}
    assert required_sheets.issubset(workbook.sheetnames)
    formulas = [workbook["Formula Lab"].cell(row=row, column=3).value for row in range(2, 8)]
    assert all(isinstance(value, str) and value.startswith("=") for value in formulas)
    assert any("XLOOKUP" in value for value in formulas)
    assert any("INDEX" in value and "MATCH" in value for value in formulas)
    assert any("SUMIFS" in value for value in formulas)
    assert any("COUNTIFS" in value for value in formulas)
    assert any("UNIQUE" in value for value in formulas)
    assert len(workbook["Raw Data"].conditional_formatting) >= 1
    assert "RawDataTable" in workbook["Raw Data"].tables
    assert "MISExceptionsTable" in workbook["Exceptions"].tables
    assert workbook["Reconciliation"]["A2"].value == "Control"
    assert workbook["MIS Control"]["A1"].value == "MIS Automation Control Center"
    assert workbook["MIS Automation"]["A1"].value == "MIS Automation Plan"
    assert workbook["Dashboard"]._charts

    validation = validate_powerbi_project_package(result["powerbi_bytes"])
    assert validation["valid"] is True
    assert validation["table_count"] >= 3
    assert validation["dimension_count"] >= 2
    assert validation["relationship_count"] >= 2
    assert validation["measure_count"] >= 5
    assert validation["time_intelligence_measure_count"] >= 3
    assert validation["rls_role_count"] == 1
    assert validation["slicer_count"] >= 1
    assert {"Drillthrough", "Tooltip"}.issubset(validation["page_types"])
    contract = validation["model_contract"]
    assert contract["model_type"] == "star_schema"
    assert contract["fact_table"]["grain"] == "one row per cleaned source record"
    assert contract["date_table"] == "DimDate"
    assert contract["rls"]["mode"] == "fail_closed_template"

    with ZipFile(BytesIO(result["powerbi_bytes"])) as package:
        assert {"MODEL_DESIGN.md", "RLS_CONFIGURATION.md", "model_contract.json"}.issubset(package.namelist())
        model_name = next(name for name in package.namelist() if name.endswith("/model.bim"))
        model = json.loads(package.read(model_name))["model"]
        assert all(item["crossFilteringBehavior"] == "oneDirection" for item in model["relationships"])
        assert model["roles"][0]["tablePermissions"][0]["filterExpression"] == "FALSE()"


def test_business_analysis_contract_and_professional_endpoints(client):
    frame = build_project_fixture("superstore_sales_dashboard", rows=48)
    analysis = build_business_analysis(frame, dataset_name="flagship_sales")
    assert analysis["narrative_contract"] == ["problem", "analysis", "evidence", "insights", "recommendation"]
    assert analysis["evidence"]
    assert analysis["insights"]
    assert analysis["recommendation"]["action"]
    assert analysis["metric_to_monitor_next"]
    assert analysis["confidence"] in {"medium", "high"}
    evidence_ids = {item["id"] for item in analysis["evidence"]}
    assert set(analysis["recommendation"]["evidence_ids"]).issubset(evidence_ids)

    context_limited = build_business_analysis(
        pd.DataFrame({"year": [2023, 2024, 2025], "technical_label": ["a", "b", "c"]}),
        dataset_name="non_business_source",
    )
    assert context_limited["status"] == "NEEDS_BUSINESS_CONTEXT"
    assert context_limited["insights"]
    assert context_limited["confidence"] == "low"
    assert "not a business-performance conclusion" in context_limited["insights"][0]["interpretation"]

    capabilities = client.get("/api/v1/platform/professional-capabilities")
    assert capabilities.status_code == 200
    assert capabilities.json()["status"] == "PASS"
    assert {item["id"] for item in capabilities.json()["domains"]} == {"sql", "excel", "power_bi", "business_analysis", "end_to_end"}

    catalog = client.get("/api/v1/project-catalog").json()
    assert catalog["flagship_count"] == 3
    assert catalog["flagship_project_ids"] == ["superstore_sales_dashboard", "customer_rfm_segmentation", "supply_chain_inventory"]
    assert all(item["portfolio_tier"] == "flagship" for item in catalog["flagship_projects"])
    assert all(len(item["end_to_end_stages"]) == 8 for item in catalog["flagship_projects"])
